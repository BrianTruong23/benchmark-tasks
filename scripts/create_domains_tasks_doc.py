from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/tht0021/Downloads/benchmark-tasks")
OUT_DIR = ROOT / "Domains & Tasks"
OUT_PATH = OUT_DIR / "Domains and Benchmark Tasks.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)
    # Override any border inherited from the built-in Title style.
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between", "bar"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    p_pr.append(borders)


def add_cell_text(cell, text, bold_lead=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(9.3)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12

    title_style = styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(7)

    for style_name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.add_run("Domains and Benchmark Tasks")
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run("Current shared canvas coverage and recommended next domain")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(64, 64, 64)

    intro = doc.add_paragraph()
    intro.add_run(
        "This summary maps the requested software domains to the benchmarks in the current local catalog. "
        "The catalog contains 10 benchmarks and 2,934 tasks. The shared canvas review file contains 251 proposed candidates; "
        "those candidates still require acceptance. Research checked on 7 October 2026 also identified ScratchWorld, which is not yet in the local catalog."
    )

    doc.add_heading("Current domain coverage", level=1)

    rows = [
        (
            "Presentation and PowerPoint",
            "PPT-Eval; GUIDE; DeskCraft; OSWorld; ParaGUIBench",
            "Microsoft PowerPoint Online; Microsoft PowerPoint; Google Slides; LibreOffice Impress",
            "Covered. PPT-Eval has 120 catalog tasks and was excluded from the candidate file because it was already reviewed. Additional presentation candidates come from GUIDE, DeskCraft, and OSWorld.",
        ),
        (
            "Microsoft Office and productivity suites",
            "GUIDE; PPT-Eval; DeskCraft; OSWorld; ParaGUIBench",
            "PowerPoint; Excel; Google Slides and Sheets; LibreOffice Writer, Calc, and Impress",
            "Covered across presentations, spreadsheets, and documents. GUIDE supplies direct Microsoft PowerPoint and Excel conditions; the other benchmarks add web and LibreOffice equivalents.",
        ),
        (
            "Scratch programming and entertainment",
            "ScratchWorld newly identified; not yet imported",
            "MIT Scratch",
            "Benchmark available. ScratchWorld was released in February 2026 with 83 curated tasks across Create, Debug, Extend, and Compute. It supports primitive GUI actions and higher-level composite actions. Add it to the catalog and then review its tasks for shared canvas suitability.",
        ),
        (
            "UI and graphic design",
            "CANVAS; GUIDE",
            "Figma; Canva",
            "Covered. CANVAS contributes 598 Figma tasks and 40 shared canvas candidates. GUIDE adds Figma and Canva tasks. CANVAS is the benchmark name; Canva is the application. The Sketch application is not represented.",
        ),
        (
            "3D and CAD design",
            "DeskCraft; CADWorld; AutoCAD Bench",
            "Blender; FreeCAD; Autodesk AutoCAD",
            "Covered. Current candidates include 17 from CADWorld, 6 Blender tasks from DeskCraft, and 1 from AutoCAD Bench.",
        ),
        (
            "Video editing and post production",
            "CutVerse; DeskCraft; GUIDE",
            "Premiere Pro; After Effects; DaVinci Resolve; JianYing; Kdenlive; CapCut; Keling; ComfyUI",
            "Covered. CutVerse contributes 10 shared canvas candidates, DeskCraft contributes 9 Kdenlive candidates, and GUIDE contributes Premiere Pro and CapCut conditions.",
        ),
        (
            "Image editing illustration and compositing",
            "PSBench; DeskCraft; GUIDE; OSWorld; CutVerse",
            "Adobe Photoshop; GIMP; Inkscape",
            "Recommended next category. About 74 current candidates involve Photoshop, GIMP, or Inkscape. Layers, masks, visual regions, typography, and effects provide natural concurrent ownership while requiring one coherent image.",
        ),
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [1.42, 1.65, 1.68, 3.01]
    headers = ["Domain", "Benchmarks", "Software", "Coverage and next action"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, cell in enumerate(header.cells):
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(headers[idx])
        run.bold = True
        run.font.name = "Aptos"
        run.font.size = Pt(9.3)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for col_idx, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_width(cell, widths[col_idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx % 2 == 0:
                set_cell_shading(cell, "F3F6F9")
            add_cell_text(cell, value)
            if col_idx == 0:
                cell.paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    doc.add_heading("Scratch benchmark finding", level=1)
    p = doc.add_paragraph()
    p.add_run("Scratch does have a benchmark as of 2026. ").bold = True
    p.add_run(
        "ScratchWorld evaluates multimodal GUI agents that construct Scratch programs through the interface. "
        "Its 83 tasks cover creating projects, debugging, extending existing projects, and computation. "
        "The correct status is therefore available externally but not yet included in this project, rather than no benchmark exists."
    )

    doc.add_heading("Recommended next step", level=1)
    p = doc.add_paragraph()
    p.add_run("Add image editing illustration and compositing as the next formal category. ").bold = True
    p.add_run(
        "It has the deepest ungrouped shared canvas coverage in the current files. After that, import ScratchWorld and assess which Scratch tasks can be divided by sprites, scripts, scenes, or functional modules without losing project-wide behavior."
    )

    doc.add_heading("Sources", level=1)
    source_entries = [
        ("Local benchmark task catalog", str(ROOT / "benchmark_tasks.json")),
        ("Local shared canvas candidate list", str(ROOT / "shared_canvas_candidates.json")),
        ("ScratchWorld paper", "https://arxiv.org/abs/2602.10814"),
        ("ScratchWorld repository", "https://github.com/astarforbae/ScratchWorld"),
        ("MIT Scratch", "https://scratch.mit.edu/"),
    ]
    for label, url in source_entries:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        if url.startswith("http"):
            add_hyperlink(p, label, url)
        else:
            p.add_run(f"{label}: {url}")

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    core = doc.core_properties
    core.title = "Domains and Benchmark Tasks"
    core.subject = "Shared canvas benchmark domain coverage"
    core.author = "Benchmark Tasks Project"
    core.keywords = "benchmarks, shared canvas, ScratchWorld, GUI agents"

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_document()
