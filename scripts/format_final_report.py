#!/usr/bin/env python3
"""Create a polished final report from the user-supplied updated DOCX.

The source content and embedded images are retained. This script normalizes the
section hierarchy, typography, task callouts, tables, spacing, header, and page
footer and writes a separate final-report file.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path("/Users/tht0021/Downloads/Domains & Tasks (Updated) (1).docx")
OUTPUT = Path("/Users/tht0021/Downloads/benchmark-tasks/Benchmark_Landscape_Final.docx")

NAVY = "16324F"
BLUE = "2E5F8A"
PALE_BLUE = "EAF1F7"
PALE_GOLD = "FFF5DC"
LIGHT = "F5F7F9"
MID = "D5DEE7"
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(91, 103, 116)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=90, start=95, bottom=90, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color=MID, size="4"):
    tbl_pr = table._tbl.tblPr
    box = tbl_pr.find(qn("w:tblBorders"))
    if box is None:
        box = OxmlElement("w:tblBorders")
        tbl_pr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = box.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            box.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def paragraph_rule(paragraph, color=BLUE, size="14", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def add_custom_styles(doc):
    styles = doc.styles

    task = styles.add_style("Task title", WD_STYLE_TYPE.PARAGRAPH) if "Task title" not in styles else styles["Task title"]
    task.font.name = "Aptos"
    task.font.size = Pt(10.2)
    task.font.bold = True
    task.font.color.rgb = DARK
    task.paragraph_format.space_before = Pt(7)
    task.paragraph_format.space_after = Pt(4)
    task.paragraph_format.keep_with_next = True

    label = styles.add_style("Report label", WD_STYLE_TYPE.PARAGRAPH) if "Report label" not in styles else styles["Report label"]
    label.font.name = "Aptos"
    label.font.size = Pt(9.2)
    label.font.bold = True
    label.font.color.rgb = RGBColor(46, 95, 138)
    label.paragraph_format.space_before = Pt(5)
    label.paragraph_format.space_after = Pt(2)
    label.paragraph_format.keep_with_next = True

    callout = styles.add_style("Report callout", WD_STYLE_TYPE.PARAGRAPH) if "Report callout" not in styles else styles["Report callout"]
    callout.font.name = "Aptos"
    callout.font.size = Pt(9.4)
    callout.font.color.rgb = DARK
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(2)
    callout.paragraph_format.space_after = Pt(5)

    caption = styles.add_style("Report caption", WD_STYLE_TYPE.PARAGRAPH) if "Report caption" not in styles else styles["Report caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.keep_with_next = True


def format_table(table, widths, font_size):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    repeat_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        prevent_split(row)
        for c_idx, cell in enumerate(row.cells):
            set_width(cell, widths[c_idx])
            margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                shade(cell, NAVY)
            elif r_idx % 2 == 0:
                shade(cell, LIGHT)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.02
                if r_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
        if r_idx > 0 and row.cells[0].paragraphs[0].runs:
            row.cells[0].paragraphs[0].runs[0].bold = True


def build():
    doc = Document(SOURCE)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 10.8, NAVY)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    list_style = styles["List Bullet"]
    list_style.font.name = "Aptos"
    list_style.font.size = Pt(9.8)
    list_style.font.color.rgb = DARK
    list_style.paragraph_format.left_indent = Inches(0.28)
    list_style.paragraph_format.first_line_indent = Inches(-0.16)
    list_style.paragraph_format.space_after = Pt(4)

    add_custom_styles(doc)

    # Remove purely empty spacer paragraphs while retaining paragraphs that hold images.
    for p in list(doc.paragraphs):
        if not p.text.strip() and not p._p.xpath(".//w:drawing"):
            remove_paragraph(p)

    paragraphs = doc.paragraphs
    title = next(p for p in paragraphs if p.text.strip() == "Computer-Use Benchmark Landscape")
    title.style = styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(2)
    for run in title.runs:
        run.font.name = "Aptos Display"
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph_rule(title, color=BLUE, size="18", space="4")

    subtitle = next(p for p in paragraphs if p.text.strip() == "Domains, Software, and Shared-Canvas Suitability")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_after = Pt(16)
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = MUTED

    renames = {
        "4. Top shared-canvas tasks per benchmark": ("4. Top shared-canvas tasks per benchmark", "Heading 1"),
        "1. PPT-Eval: PPT-online:": ("4.1 PPT-Eval — PowerPoint Online", "Heading 2"),
        "2. Canvas: Figma": ("4.2 CANVAS — Figma", "Heading 2"),
        "3. PSBench: Photoshop": ("4.3 PSBench — Adobe Photoshop", "Heading 2"),
        "5. ScratchWorld": ("4.4 ScratchWorld — MIT Scratch", "Heading 2"),
        "6. ProSoft Arena": ("4.5 ProSoftArena — Professional software", "Heading 2"),
        "7. OSWORLD": ("4.6 OSWorld — LibreOffice", "Heading 2"),
        "5. Main Concerns": ("5. Main concerns", "Heading 1"),
        "6. Next Steps": ("6. Next steps", "Heading 1"),
        "1.  Microsoft Excel:": ("Microsoft Excel", "Heading 3"),
        "2.  ILLUSTRATOR": ("Adobe Illustrator", "Heading 3"),
        "A. LibreOfficeCal": ("LibreOffice Calc", "Heading 3"),
        "B. LibreOfficeImpress": ("LibreOffice Impress", "Heading 3"),
    }
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in renames:
            new_text, style_name = renames[text]
            p.text = new_text
            p.style = styles[style_name]

    main_headings = {"1. Executive summary", "2. Scale of the landscape", "3. Benchmarks at a glance"}
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in main_headings:
            p.style = styles["Heading 1"]
        if p.style.name == "Heading 1":
            paragraph_rule(p, color=MID, size="6", space="2")

    # Normalize task blocks and shared-canvas callouts.
    for p in doc.paragraphs:
        text = p.text.strip()
        has_image = bool(p._p.xpath(".//w:drawing"))
        if has_image:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
            continue
        if text.startswith("Example Task ") or text.startswith("Task: ") or text.startswith("A. Build ") or text.startswith("A. Design "):
            p.style = styles["Task title"]
        elif text[:3] in {"1. ", "2. "} and p.style.name == "normal":
            p.style = styles["Task title"]
        elif text in {"Why this is a shared-canvas task:", "=> Why this is a shared-canvas task:", "Agent assignments:"}:
            if text.startswith("=>"):
                p.text = "Why this is a shared-canvas task:"
            p.style = styles["Report label"]
        elif text.startswith("=>"):
            p.text = text[2:].strip()
            p.style = styles["Report callout"]
        elif text.startswith("Agent A:") or text.startswith("Agent B:") or text.startswith("Agent C:"):
            p.style = styles["List Bullet"]
        elif text.startswith("Starting") and "Target" in text:
            p.text = "Starting artifact                                     Target artifact"
            p.style = styles["Report caption"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text.startswith("•"):
            p.text = text.lstrip("•").strip()
            p.style = styles["List Bullet"]

    # Give rationale paragraphs a subtle callout treatment.
    previous_was_label = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if previous_was_label and text and not text.startswith("Agent "):
            p.style = styles["Report callout"]
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_BLUE)
            p_pr.append(shd)
        previous_was_label = text == "Why this is a shared-canvas task:"

    # Style images and keep each image row together.
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_together = True

    format_table(doc.tables[0], [0.82, 0.85, 0.43, 1.18, 2.82, 1.15], 7.2)
    format_table(doc.tables[1], [2.05, 5.2], 9)

    # Header and footer.
    header = section.header.paragraphs[0]
    header.text = "COMPUTER-USE BENCHMARK LANDSCAPE"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(2)
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = MUTED
    paragraph_rule(header, color=MID, size="4", space="2")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Final report  •  ")
    set_field(footer, "PAGE")
    footer.add_run(" of ")
    set_field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    core = doc.core_properties
    core.title = "Computer-Use Benchmark Landscape"
    core.subject = "Domains, software, and shared-canvas suitability"
    core.author = "Benchmark Tasks Project"
    core.keywords = "computer-use benchmarks, shared canvas, multi-agent"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
